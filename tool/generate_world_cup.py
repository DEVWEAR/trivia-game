#!/usr/bin/env python3
import json, os, time, urllib.parse, urllib.request
from pathlib import Path
from openpyxl import load_workbook

root=Path(__file__).resolve().parents[1]
src=root/'tool/imports/world_cup_question_bank_300_fixed.xlsx'
assert src.exists(), 'Upload the provided workbook to tool/imports/world_cup_question_bank_300_fixed.xlsx'
wb=load_workbook(src,read_only=True,data_only=True)
rows=list(wb['Question Bank'].iter_rows(min_row=2,max_col=7,values_only=True))
assert len(rows)==300
assert len({r[0] for r in rows})==300
assert len({str(r[5]).strip() for r in rows})==300
assert {p:sum(r[3]==p for r in rows) for p in (200,400,600)}=={200:100,400:100,600:100}
cache_path=root/'tool/imports/world_cup_en_cache.json'
cache=json.loads(cache_path.read_text()) if cache_path.exists() else {}
def translate(s):
    params=urllib.parse.urlencode({'client':'gtx','sl':'ar','tl':'en','dt':'t','q':str(s)})
    for attempt in range(7):
        try:
            req=urllib.request.Request('https://translate.googleapis.com/translate_a/single?'+params,headers={'User-Agent':'Mozilla/5.0'})
            with urllib.request.urlopen(req,timeout=30) as res: data=json.load(res)
            result=''.join(t[0] for t in data[0] if t[0])
            if result.strip():return result.strip()
        except Exception:
            time.sleep(min(2**attempt,25))
    raise RuntimeError('Translation failed: '+str(s)[:70])
def dart(s):
    return "'"+str(s).replace('\\','\\\\').replace("'","\\'").replace('$','\\$').replace('\n',' ')+"'"
out=["import '../question_model.dart';","","// 300 World Cup questions from the supplied workbook; English machine translations require editorial review.","final worldCupQuestions = <TriviaQuestion>["]
for i,r in enumerate(rows,1):
    qid,cat,difficulty,points,topic,question,answer=r
    if qid not in cache:
        cache[qid]={'q':translate(question),'a':translate(answer)}
        if i%10==0:
            cache_path.write_text(json.dumps(cache,ensure_ascii=False,indent=2))
    en=cache[qid]
    args=[
        'id:'+dart(qid),'categoryId:\'world_cup\'',
        'difficulty:QuestionDifficulty.'+{200:'easy200',400:'medium400',600:'hard600'}[points],
        'questionAr:'+dart(question),'questionEn:'+dart(en['q']),
        'answerAr:'+dart(answer),'answerEn:'+dart(en['a']),
        "sourceName:'User-provided World Cup workbook (editorial review pending)'",
        "sourceUrl:'https://www.fifa.com/en/tournaments/mens/worldcup'",
        'lastVerified:DateTime(2026,10,10)',
        'factKey:'+dart('worldcup:'+qid)
    ]
    out.append('  TriviaQuestion('+','.join(args)+'),')
cache_path.write_text(json.dumps(cache,ensure_ascii=False,indent=2))
out.append('];')
dest=root/'lib/data/questions/world_cup_questions.dart'
dest.write_text('\n'.join(out)+'\n')
assert dest.read_text().count('TriviaQuestion(')==300
print('Generated 300 bilingual questions')
